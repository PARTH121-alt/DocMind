"""Text extraction for PDF, Office, tabular, plain-text and image formats.

Extracted text is returned page by page so citations can point at a specific
page, slide, or sheet. OCR is invoked only when a page yields too little text
to be a real text layer (i.e. a scan).
"""

from __future__ import annotations

import io
import json
import logging
import re
from pathlib import Path

from app.core.config import settings
from app.services.ai import hf_api
from app.services.rag.types import ExtractedDocument, Page

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".csv", ".xlsx", ".pptx",
    ".md", ".markdown", ".json", ".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff",
}

# A text layer below this many characters per page is assumed to be a scan.
OCR_CHAR_THRESHOLD = 120


class ExtractionError(RuntimeError):
    pass


def _needs_ocr(text: str) -> bool:
    return len(re.sub(r"\s+", "", text)) < OCR_CHAR_THRESHOLD


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError("PyMuPDF is required for PDF extraction") from exc

    doc = fitz.open(path)
    meta = {k: v for k, v in (doc.metadata or {}).items() if v and k != "format"}
    pages: list[Page] = []

    for i, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        used_ocr = False

        if _needs_ocr(text):
            ocr_text = _ocr_pdf_page(page, i)
            if ocr_text:
                text = ocr_text
                used_ocr = True

        pages.append(Page(number=i, text=text, section=_pdf_section(page), used_ocr=used_ocr))

    doc.close()
    return ExtractedDocument(filename=path.name, pages=pages, metadata=meta)


def _pdf_section(page) -> str | None:
    """Use the largest-font line on the page as a heading hint."""
    try:
        blocks = page.get_text("dict").get("blocks", [])
    except Exception:
        return None
    best, best_size = None, 0.0
    for block in blocks:
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                size = span.get("size", 0)
                if text and size > best_size and len(text) < 120:
                    best, best_size = text, size
    return best


def _ocr_pdf_page(page, page_number: int) -> str:
    """Rasterize and OCR a scanned PDF page."""
    text = _ocr_image_bytes(page.get_pixmap(dpi=200).tobytes("png"))
    if text:
        logger.info("OCR applied to PDF page %s", page_number)
    return text


# --------------------------------------------------------------------------
# Office formats
# --------------------------------------------------------------------------
def _extract_docx(path: Path) -> ExtractedDocument:
    import docx

    doc = docx.Document(str(path))
    parts: list[str] = []
    section = None
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name.lower().startswith(("heading", "title")):
            section = text
            parts.append(f"\n## {text}")
        else:
            parts.append(text)

    for t_idx, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            parts.append(f"\n[TABLE {t_idx}]\n" + "\n".join(rows))

    body = "\n".join(parts)
    return ExtractedDocument(
        filename=path.name,
        pages=[Page(number=1, text=body, section=section)],
        metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
    )


def _extract_pptx(path: Path) -> ExtractedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages: list[Page] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    lines.append(t)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        title = slide.shapes.title.text.strip() if slide.shapes.title else None
        pages.append(Page(number=i, text="\n".join(lines), section=title))
    return ExtractedDocument(
        filename=path.name, pages=pages, metadata={"slides": len(prs.slides)}
    )


def _extract_xlsx(path: Path) -> ExtractedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    pages: list[Page] = []
    for idx, ws in enumerate(wb.worksheets, start=1):
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        pages.append(Page(number=idx, text="\n".join(rows), section=ws.title))
    wb.close()
    return ExtractedDocument(
        filename=path.name, pages=pages, metadata={"sheets": len(pages)}
    )


def _extract_csv(path: Path) -> ExtractedDocument:
    import csv as csv_module

    with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        rows = []
        reader = csv_module.reader(fh)
        for i, row in enumerate(reader):
            rows.append(" | ".join(row))
            if i > 20000:  # guard against pathological files
                break
    return ExtractedDocument(
        filename=path.name, pages=[Page(number=1, text="\n".join(rows))], metadata={"type": "csv"}
    )


def _extract_json(path: Path) -> ExtractedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(raw)
        text = json.dumps(data, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        text = raw
    return ExtractedDocument(filename=path.name, pages=[Page(number=1, text=text)], metadata={"type": "json"})


def _extract_text(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ExtractedDocument(filename=path.name, pages=[Page(number=1, text=text)])


# --------------------------------------------------------------------------
# Images
# --------------------------------------------------------------------------
def _extract_image(path: Path) -> ExtractedDocument:
    text = _ocr_image_bytes(path.read_bytes())
    if not text:
        raise ExtractionError(
            "No text could be read from this image. Configure OCR (set HF_TOKEN for "
            "the TrOCR model) and try again."
        )
    return ExtractedDocument(
        filename=path.name,
        pages=[Page(number=1, text=text, used_ocr=True)],
        metadata={"type": "image"},
    )


def _ocr_image_bytes(data: bytes) -> str:
    """OCR an image using the HF TrOCR model."""
    if not hf_api.is_configured():
        logger.warning("Image/scanned page needs OCR but HF_TOKEN is not configured.")
        return ""
    try:
        from PIL import Image

        image = Image.open(io.BytesIO(data)).convert("RGB")
        client = hf_api._client()
        raw = client.image_to_text(image, model=settings.hf_ocr_model)
        return raw if isinstance(raw, str) else str(raw)
    except Exception as exc:
        logger.warning("OCR failed: %s", exc)
        return ""


# --------------------------------------------------------------------------
# Dispatch
# --------------------------------------------------------------------------
_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".csv": _extract_csv,
    ".json": _extract_json,
    ".txt": _extract_text,
    ".md": _extract_text,
    ".markdown": _extract_text,
    ".png": _extract_image,
    ".jpg": _extract_image,
    ".jpeg": _extract_image,
    ".webp": _extract_image,
    ".gif": _extract_image,
    ".tif": _extract_image,
    ".tiff": _extract_image,
}


def extract(path: Path) -> ExtractedDocument:
    """Extract text from a file, dispatching on its extension."""
    ext = path.suffix.lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        raise ExtractionError(f"Unsupported file type: {ext or 'unknown'}")
    return fn(path)
